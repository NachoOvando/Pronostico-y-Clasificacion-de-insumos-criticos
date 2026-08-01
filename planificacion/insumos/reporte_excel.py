"""Export del resultado del K-Means a Excel, con formato de informe.

No calcula nada: recibe las dos tablas ya armadas y las escribe con estilo,
incrustando los PNG de la Sección 7.
"""

import logging
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from planificacion import estilos
from planificacion.config import InsumosConfig

logger = logging.getLogger(__name__)

_FINA = Side(style='thin', color='CCCCCC')
_GRUESA = Side(style='medium', color=estilos.hex_openpyxl(estilos.AZUL))
_BORDE_DATO = Border(left=_FINA, right=_FINA, top=_FINA, bottom=_FINA)
_BORDE_ENCABEZADO = Border(left=_GRUESA, right=_GRUESA, top=_GRUESA, bottom=_GRUESA)

_FILA_PAR = 'F2F2F2'
_FILA_IMPAR = 'FFFFFF'


def _relleno(color_hex: str) -> PatternFill:
    return PatternFill('solid', fgColor=estilos.hex_openpyxl(color_hex))


def _titulo(ws: Worksheet, fila: int, n_cols: int, texto: str,
            fondo: str = estilos.AZUL, color: str = 'FFFFFF', tam: int = 13) -> None:
    ws.merge_cells(f'A{fila}:{get_column_letter(n_cols)}{fila}')
    celda = ws.cell(row=fila, column=1, value=texto)
    celda.font = Font(bold=True, size=tam, color=estilos.hex_openpyxl(color), name='Arial')
    celda.fill = _relleno(fondo)
    celda.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[fila].height = 28


def _encabezado(ws: Worksheet, fila: int, etiquetas: list[str]) -> None:
    for col, etiqueta in enumerate(etiquetas, 1):
        celda = ws.cell(row=fila, column=col, value=etiqueta)
        celda.font = Font(bold=True, size=9, color='FFFFFF', name='Arial')
        celda.fill = _relleno(estilos.AZUL)
        celda.border = _BORDE_ENCABEZADO
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[fila].height = 30


def _celda(ws: Worksheet, fila: int, col: int, valor, fondo: str = 'FFFFFF',
           negrita: bool = False, formato: str | None = None,
           color: str = '000000', alineacion: str = 'center') -> None:
    celda = ws.cell(row=fila, column=col, value=valor)
    celda.fill = _relleno(fondo)
    celda.border = _BORDE_DATO
    celda.font = Font(bold=negrita, size=9, name='Arial', color=estilos.hex_openpyxl(color))
    celda.alignment = Alignment(horizontal=alineacion, vertical='center')
    if formato:
        celda.number_format = formato


def _anchos(ws: Worksheet, columnas: str, anchos: list[int]) -> None:
    for col, ancho in zip(columnas, anchos):
        ws.column_dimensions[col].width = ancho


def _insertar_imagen(ws: Worksheet, path: Path, ancla: str,
                     ancho: int, alto: int) -> None:
    try:
        imagen = XLImage(path)
    except FileNotFoundError:
        logger.warning("Falta %s: ejecutá la celda del gráfico antes de exportar.", path)
        return
    imagen.width, imagen.height = ancho, alto
    ws.add_image(imagen, ancla)


def _hoja_familias(ws: Worksheet, df_resultado: pd.DataFrame, cfg: InsumosConfig) -> None:
    ws.title = 'Familias de Compra'
    _titulo(ws, 1, 11,
            f'INSUMOS CRÍTICOS POR FAMILIA DE COMPRA — K-Means (K={cfg.k_clusters})')
    _titulo(ws, 2, 11,
            'Features: Volumen relativo por UM  ×  Alcance productivo (%)  ×  Lead Time',
            fondo=estilos.AZUL_MEDIO, tam=9)
    _encabezado(ws, 3, ['#', 'Familia de Compra', 'UM', 'Consumo Total',
                        'Variantes (talles)', 'N Artículos', 'Alcance (%)',
                        'Vol. relativo (%)', f'Lead Time ({cfg.leadtime_unidad})',
                        'Proveedor', 'Criticidad'])

    for i, (_, fila) in enumerate(df_resultado.iterrows(), 1):
        n_fila = i + 3
        fondo = _FILA_PAR if i % 2 == 0 else _FILA_IMPAR
        criticidad = fila['Criticidad']
        valores = [
            (i, None, 'center'),
            (fila['Familia'], None, 'left'),
            (fila['UM'], None, 'center'),
            (fila['Consumo_Total'], '#,##0.0', 'center'),
            (fila['N_SKUs_hijo'], None, 'center'),
            (fila['N_Articulos'], None, 'center'),
            (fila['Alcance_pct'], '0.0', 'center'),
            (fila['Vol_norm'], '0.0', 'center'),
            (fila['Lead_Time_dias'], '0', 'center'),
            (fila['Proveedor'], None, 'left'),
        ]
        for col, (valor, formato, alineacion) in enumerate(valores, 1):
            _celda(ws, n_fila, col, valor, fondo=fondo, formato=formato,
                   alineacion=alineacion)
        _celda(ws, n_fila, 11, criticidad,
               fondo=estilos.COLOR_CRIT[criticidad], negrita=True,
               color=estilos.COLOR_TEXTO_CRIT[criticidad])

    _anchos(ws, 'ABCDEFGHIJK', [5, 38, 6, 16, 12, 12, 10, 14, 12, 18, 14])
    ws.freeze_panes = 'A4'

    n = len(df_resultado)
    _insertar_imagen(ws, cfg.grafico_dispersion, f'A{n + 6}', 780, 490)
    _insertar_imagen(ws, cfg.grafico_paralelas, f'F{n + 32}', 480, 340)
    _insertar_imagen(ws, cfg.grafico_heatmap, f'A{n + 56}',
                     520, int(520 * (0.42 * n + 2.2) / 9.5))


def _hoja_skus(ws: Worksheet, df_sku: pd.DataFrame) -> None:
    _titulo(ws, 1, 7, 'DETALLE — SKUs Hijo por Familia de Compra')
    _titulo(ws, 2, 7,
            'Variantes de talle que componen cada familia · '
            'Útil para emisión de órdenes de compra',
            fondo=estilos.AZUL_MEDIO, tam=9)
    _encabezado(ws, 3, ['Criticidad', 'Familia', 'UM', 'SKU Hijo (con talle)',
                        'Consumo SKU', 'N Artículos', 'Alcance (%)'])

    familia_anterior = None
    for i, (_, fila) in enumerate(df_sku.iterrows(), 1):
        n_fila = i + 3
        fondo = _FILA_PAR if i % 2 == 0 else _FILA_IMPAR
        criticidad = fila['Criticidad']

        # La criticidad y el nombre de familia se escriben solo en la primera
        # fila del grupo, para que el bloque se lea como una unidad de compra.
        if fila['Familia'] != familia_anterior:
            fondo_fam = estilos.COLOR_CRIT.get(criticidad, '#FFFFFF')
            color_fam = estilos.COLOR_TEXTO_CRIT.get(criticidad, '#000000')
            _celda(ws, n_fila, 1, criticidad, fondo=fondo_fam, negrita=True, color=color_fam)
            _celda(ws, n_fila, 2, fila['Familia'], fondo=fondo_fam, negrita=True,
                   color=color_fam, alineacion='left')
            familia_anterior = fila['Familia']
        else:
            _celda(ws, n_fila, 1, '', fondo=fondo)
            _celda(ws, n_fila, 2, '', fondo=fondo)

        _celda(ws, n_fila, 3, fila['UM'], fondo=fondo)
        _celda(ws, n_fila, 4, fila['Nombre_Base'], fondo=fondo, alineacion='left')
        _celda(ws, n_fila, 5, fila['Consumo_SKU'], fondo=fondo, formato='#,##0.0')
        _celda(ws, n_fila, 6, fila['N_Art'], fondo=fondo)
        _celda(ws, n_fila, 7, fila['Alcance_pct'], fondo=fondo, formato='0.0')

    _anchos(ws, 'ABCDEFG', [14, 38, 6, 38, 18, 12, 12])
    ws.freeze_panes = 'A4'


def exportar_informe_de_criticidad(df_resultado: pd.DataFrame, df_sku: pd.DataFrame,
             cfg: InsumosConfig) -> Path:
    """Genera el Excel final con las dos hojas del informe."""
    wb = Workbook()
    _hoja_familias(wb.active, df_resultado, cfg)
    _hoja_skus(wb.create_sheet('Detalle por SKU hijo'), df_sku)
    wb.save(cfg.output_path)
    return cfg.output_path
